"""Exporta exclusivamente las cinco fuentes de Actividades a Word editable."""
from pathlib import Path
from importlib.util import spec_from_file_location, module_from_spec
import json
from lxml import html, etree
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

BASE = Path(__file__).resolve().parents[1]
spec = spec_from_file_location('editorial', BASE/'scripts/render-contenido-word-nativo.py')
m = module_from_spec(spec)
spec.loader.exec_module(m)
report = []
for n in range(1, 6):
    folder = BASE/'entregables'/f'MODULO {n}'/'Actividades'
    tree = html.fromstring((folder/'actividades.html').read_text(encoding='utf-8-sig'))
    article = tree.xpath('//article')[0]
    doc = Document()
    m.styles(doc, '173F67')
    for name in ['Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    doc.styles['Title'].font.size = Pt(26)
    doc.styles['Heading 1'].font.size = Pt(16)
    doc.core_properties.title = f'Actividades del módulo {n} Gestión de Proyectos Ágiles'
    doc.core_properties.subject = 'Práctica autónoma con respuestas orientativas'
    b = m.Builder(doc, folder)
    for element in article:
        if element.tag == 'h1':
            doc.add_paragraph(element.text_content(), style='Title')
        else:
            if element.tag == 'h2' and element.text_content().startswith('Respuestas'):
                doc.add_page_break()
            b.block(element)
    for table in doc.tables:
        for row in table.rows:
            pr = row._tr.get_or_add_trPr()
            pr.append(OxmlElement('w:cantSplit'))
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                borders = OxmlElement('w:tcBorders')
                for side in ['top', 'left', 'bottom', 'right']:
                    border = OxmlElement('w:'+side)
                    for key, value in [('val','single'),('sz','4'),('color','D9D9D9')]:
                        border.set(qn('w:'+key), value)
                    borders.append(border)
                cell._tc.get_or_add_tcPr().append(borders)
    m.finalize(doc)
    target = folder/'Actividades.docx'
    doc.save(target)
    # Verifica que todos los bloques de la fuente estén presentes en el Word.
    xml = doc.element
    word_text = ' '.join(xml.xpath('//w:t/text()'))
    normalize = lambda s: ''.join(s.split())
    missing = []
    for e in article.xpath('.//h1|.//h2|.//h3|.//p|.//li|.//td|.//th'):
        # El exportador puede separar párrafos largos en sus oraciones.
        if normalize(e.text_content()) not in normalize(word_text):
            missing.append(e.text_content()[:80])
    assert not missing, (n, missing)
    assert '\ufffd' not in word_text
    count = len(article.xpath('./h2[starts-with(text(),"Actividad ")]'))
    report.append({'modulo': n, 'actividades': count, 'palabras': len(word_text.split()),
                   'tablas': len(doc.tables), 'enlaces': len(article.xpath('.//a')),
                   'texto_completo': True, 'paginacion_nativa': 'pendiente por falta de LibreOffice'})
out = BASE/'.qa/actividades-2026-09-19'
out.mkdir(parents=True, exist_ok=True)
(out/'exportacion.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(report, ensure_ascii=False, indent=2))
