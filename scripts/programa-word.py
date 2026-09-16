from pathlib import Path
from importlib.util import spec_from_file_location,module_from_spec
from docx import Document
BASE=Path(__file__).resolve().parents[1]
spec=spec_from_file_location('exportador',BASE/'scripts/render-contenido-word-nativo.py'); m=module_from_spec(spec); spec.loader.exec_module(m)
source=(BASE/'PROGRAMA-CONTENIDOS.txt').read_text(encoding='utf-8-sig').split('NOTAS INTERNAS DE ELABORACIÓN')[0]
doc=Document(); m.styles(doc,'173F67'); doc.core_properties.title='Gestión de Proyectos Ágiles Contenidos del programa'
doc.add_paragraph('Gestión de Proyectos Ágiles',style='Title')
doc.add_paragraph('Contenidos del programa',style='Subtitle')
doc.add_paragraph('Temario organizado a partir de los contenidos mínimos proporcionados para la materia. El recorrido comienza por los fundamentos, desarrolla los marcos de trabajo y su aplicación al producto, y finaliza con seguimiento, colaboración y casos de estudio.')
for line in source.splitlines()[3:]:
    if line.startswith('Módulo '): doc.add_paragraph(line,style='Heading 1')
    elif line.startswith('›'): doc.add_paragraph(line[1:].strip())
doc.add_paragraph('Este documento contiene el apartado de contenidos. La carga horaria, las correlatividades y las condiciones de evaluación y aprobación deberán incorporarse cuando se cuente con los requisitos institucionales.',style='Caption')
m.finalize(doc); doc.save(BASE/'Programa - contenidos.docx')
print('Programa - contenidos.docx generado.')
