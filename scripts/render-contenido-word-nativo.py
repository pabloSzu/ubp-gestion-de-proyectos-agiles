"""Word académico editable: estilos, listas, tablas, enlaces y recursos del HTML."""
from pathlib import Path
import re,json
from lxml import html
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image

BASE=Path(__file__).resolve().parents[1]
PALETTE=['173F67','12655D','5D3E83','815B19','8C3838']

def shade(p,fill):
    pr=p._p.get_or_add_pPr(); e=OxmlElement('w:shd'); e.set(qn('w:fill'),fill); pr.append(e)
def hyperlink(p,label,url,bold=False,italic=False):
    h=OxmlElement('w:hyperlink')
    if url.startswith('#'): h.set(qn('w:anchor'),url[1:])
    else: h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r'); pr=OxmlElement('w:rPr')
    style=OxmlElement('w:rStyle'); style.set(qn('w:val'),'Hyperlink'); pr.append(style)
    if bold: pr.append(OxmlElement('w:b'))
    if italic: pr.append(OxmlElement('w:i'))
    r.append(pr); t=OxmlElement('w:t'); t.set(qn('xml:space'),'preserve'); t.text=label; r.append(t); h.append(r); p._p.append(h)
def tokens(e,bold=False,italic=False,url=None):
    if e.tag in ('strong','b'): bold=True
    if e.tag in ('em','i'): italic=True
    if e.tag=='a': url=e.get('href')
    out=[]
    if e.text: out.append((e.text,bold,italic,url))
    for c in e:
        if c.tag=='br': out.append(('\n',bold,italic,url))
        elif c.tag not in ('script','style','img'): out.extend(tokens(c,bold,italic,url))
        if c.tail: out.append((c.tail,bold,italic,url))
    return out
def clean(ts):
    out=[]; previous_space=True
    for s,b,i,u in ts:
        s=re.sub(r'\s+',' ',s)
        if previous_space: s=s.lstrip()
        if s: out.append((s,b,i,u)); previous_space=s.endswith(' ')
    if out: out[-1]=(out[-1][0].rstrip(),*out[-1][1:])
    return out
def chunks(ts,limit=115):
    raw=''.join(t[0] for t in ts)
    if len(raw.split())<=limit: return [ts]
    boundaries=[m.end() for m in re.finditer(r'[.!?](?:[”"\)])?\s+(?=[A-ZÁÉÍÓÚÑ¿¡])',raw)]
    ends=[]; start=0
    for end in boundaries:
        if len(raw[start:end].split())>=65: ends.append(end); start=end
    if not ends: return [ts]
    ends.append(len(raw)); result=[]; start=0
    for end in ends:
        if end==start: continue
        subset=[]; pos=0
        for s,b,i,u in ts:
            lo=max(start-pos,0); hi=min(end-pos,len(s))
            if hi>lo: subset.append((s[lo:hi],b,i,u))
            pos+=len(s)
        if subset: result.append(clean(subset))
        start=end
    return result
def runs(p,ts):
    for s,b,i,u in ts:
        if u: hyperlink(p,s,u,b,i)
        else:
            r=p.add_run(s); r.bold=True if b else None; r.italic=True if i else None
            # Evitar que el tema por defecto o un visor sustituyan la tipografía.
            r.font.name='Calibri'
            if p.style.name=='Normal': r.font.size=Pt(11)
def numbering(doc,ordered,level=0):
    root=doc.part.numbering_part.element
    aid=max([int(e.get(qn('w:abstractNumId'))) for e in root.findall(qn('w:abstractNum'))]+[-1])+1
    nid=max([int(e.get(qn('w:numId'))) for e in root.findall(qn('w:num'))]+[0])+1
    abstract=OxmlElement('w:abstractNum'); abstract.set(qn('w:abstractNumId'),str(aid))
    lvl=OxmlElement('w:lvl'); lvl.set(qn('w:ilvl'),'0')
    for tag,val in [('start','1'),('numFmt','decimal' if ordered else 'bullet'),('lvlText','%1.' if ordered else '•')]:
        el=OxmlElement('w:'+tag); el.set(qn('w:val'),val); lvl.append(el)
    pr=OxmlElement('w:pPr'); ind=OxmlElement('w:ind'); ind.set(qn('w:left'),str(380+level*300)); ind.set(qn('w:hanging'),'220'); pr.append(ind); lvl.append(pr)
    abstract.append(lvl); root.append(abstract)
    num=OxmlElement('w:num'); num.set(qn('w:numId'),str(nid)); ref=OxmlElement('w:abstractNumId'); ref.set(qn('w:val'),str(aid)); num.append(ref); root.append(num)
    return nid
def list_format(p,nid):
    pr=p._p.get_or_add_pPr(); num=OxmlElement('w:numPr')
    for tag,val in [('ilvl','0'),('numId',str(nid))]:
        e=OxmlElement('w:'+tag); e.set(qn('w:val'),val); num.append(e)
    pr.append(num)
def bookmark(p,name,n):
    begin=OxmlElement('w:bookmarkStart'); begin.set(qn('w:id'),str(n)); begin.set(qn('w:name'),name)
    end=OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'),str(n)); p._p.insert(1 if p._p.find(qn('w:pPr')) is not None else 0,begin); p._p.append(end)
def styles(doc,accent):
    s=doc.sections[0]; s.page_width=Cm(21); s.page_height=Cm(29.7)
    s.top_margin=Cm(2); s.bottom_margin=Cm(2); s.left_margin=Cm(2.2); s.right_margin=Cm(2.2)
    s.header_distance=Cm(.8); s.footer_distance=Cm(.8)
    normal=doc.styles['Normal']; normal.font.name='Calibri'; normal.font.size=Pt(11)
    normal.font.color.rgb=RGBColor.from_string('172333')
    normal.paragraph_format.line_spacing=1.16; normal.paragraph_format.space_after=Pt(7)
    normal.paragraph_format.widow_control=True
    for name,size,color in [('Title',29,'000000'),('Subtitle',13,'536173'),('Heading 1',17,accent),('Heading 2',13,'12655D'),('Heading 3',11,'5D3E83')]:
        st=doc.styles[name]; st.font.name='Calibri'; st.font.size=Pt(size); st.font.color.rgb=RGBColor.from_string(color)
        st.font.bold=name.startswith('Heading'); st.paragraph_format.space_before=Pt(15 if name.startswith('Heading') else 6); st.paragraph_format.space_after=Pt(7)
        st.paragraph_format.keep_with_next=name.startswith('Heading')
    cap=doc.styles['Caption']; cap.font.name='Calibri'; cap.font.size=Pt(9); cap.font.color.rgb=RGBColor.from_string('536173')
    cap.paragraph_format.space_after=Pt(10)
    if 'Hyperlink' not in doc.styles:
        from docx.enum.style import WD_STYLE_TYPE
        doc.styles.add_style('Hyperlink',WD_STYLE_TYPE.CHARACTER)
    doc.styles['Hyperlink'].font.color.rgb=RGBColor.from_string('245E91'); doc.styles['Hyperlink'].font.underline=True
    footer=s.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=footer.add_run('Gestión de Proyectos Ágiles  •  '); r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string('536173')
    field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
    doc.core_properties.author=''; doc.core_properties.subject='Contenido de Gestión de Proyectos Ágiles'
    doc.core_properties.language='es-AR'
    for border in list(doc.styles['Title'].element.iter(qn('w:pBdr'))): border.getparent().remove(border)
    for el in doc.styles.element.iter():
        for key in list(el.attrib):
            if key.split('}')[-1] in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme','themeColor','themeTint','themeShade'):
                del el.attrib[key]
    default=doc.styles.element.find(qn('w:docDefaults'))
    if default is not None:
        rpr=default.find(qn('w:rPrDefault')).find(qn('w:rPr'))
        font=rpr.find(qn('w:rFonts')); font.set(qn('w:ascii'),'Calibri'); font.set(qn('w:hAnsi'),'Calibri')
        size=rpr.find(qn('w:sz')); size.set(qn('w:val'),'22')

def finalize(doc):
    for p in doc.paragraphs:
        if p.style.name in ('Title','Subtitle','Heading 1','Heading 2','Heading 3'):
            for r in p.runs:
                r.font.name='Calibri'; r.font.size=p.style.font.size
                r.font.color.rgb=p.style.font.color.rgb
                r.bold=True if p.style.name.startswith('Heading') else False
            p.paragraph_format.keep_with_next=True
        if p.style.name=='Title':
            p.paragraph_format.space_before=Pt(14); p.paragraph_format.space_after=Pt(14)
            pr=p._p.get_or_add_pPr()
            for border in list(pr.findall(qn('w:pBdr'))): pr.remove(border)

class Builder:
    def __init__(self,doc,folder,prefix=''):
        self.doc=doc; self.folder=folder; self.prefix=prefix; self.bookmarks=max([int(e.get(qn('w:id'))) for e in doc.element.iter(qn('w:bookmarkStart'))]+[0]); self.stats={'parrafos_divididos':0,'figuras':0,'tablas':0}
    def paragraph(self,e,style=None,fill=None,container=None):
        ts=clean(tokens(e)); parts=chunks(ts) if e.tag=='p' else [ts]
        self.stats['parrafos_divididos']+=len(parts)-1
        for part in parts:
            p=(container or self.doc).add_paragraph(style=style); runs(p,part)
            p.paragraph_format.line_spacing=1.16; p.paragraph_format.space_after=Pt(7)
            if fill:
                shade(p,fill); p.paragraph_format.left_indent=Cm(.25); p.paragraph_format.right_indent=Cm(.2)
            if style=='Caption': p.paragraph_format.keep_with_next=False
    def block(self,e,fill=None,level=0):
        tag=e.tag
        if not isinstance(tag,str): return
        if tag in ('h2','h3','h4'):
            p=self.doc.add_paragraph(style={'h2':'Heading 1','h3':'Heading 2','h4':'Heading 3'}[tag]); runs(p,clean(tokens(e)))
            for r in p.runs:
                r.font.size=p.style.font.size; r.font.color.rgb=p.style.font.color.rgb; r.bold=True
            if tag=='h2': shade(p,'EAF4FA')
            if e.get('id'):
                self.bookmarks+=1; bookmark(p,self.prefix+e.get('id'),self.bookmarks)
        elif tag=='p': self.paragraph(e,fill=fill or ('EAF4FA' if 'bloque-narrativo' in e.get('class','') else None))
        elif tag in ('ul','ol'):
            nid=numbering(self.doc,tag=='ol',level)
            for li in e:
                if li.tag!='li': continue
                p=self.doc.add_paragraph(); ts=clean(tokens(li)) if not li.xpath('./ul|./ol') else clean([(li.text or '',False,False,None)]+sum([tokens(c)+([(c.tail,False,False,None)] if c.tail else []) for c in li if c.tag not in ('ul','ol')],[]))
                runs(p,ts); list_format(p,nid)
                p.paragraph_format.line_spacing=1.16; p.paragraph_format.space_after=Pt(7)
                if fill: shade(p,fill)
                for sub in li.xpath('./ul|./ol'): self.block(sub,fill,level+1)
        elif tag=='table': self.table(e)
        elif tag=='figure':
            for im in e.xpath('.//img'):
                path=(self.folder/im.get('src')).resolve()
                if not path.is_file(): raise FileNotFoundError(path)
                with Image.open(path) as image: w,h=image.size
                width=min(16.6,8.6*w/h)
                p=self.doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.keep_with_next=True
                drawing=p.add_run().add_picture(str(path),width=Cm(width)); drawing._inline.docPr.set('descr',im.get('alt','Diagrama explicativo'))
                self.stats['figuras']+=1
            for caption in e.xpath('./figcaption'): self.paragraph(caption,'Caption')
        elif tag=='img':
            wrapper=html.Element('figure'); wrapper.append(e); self.block(wrapper)
        elif tag in ('div','blockquote','section','aside'):
            css=e.get('class',''); selected=fill
            for key,color in [('analogia','F2EDF8'),('mito','FBEEEE'),('criterio-profesional','E8F6F3'),('caso-resuelto','EAF4FA'),('epigraph','FBF2DD')]:
                if key in css: selected=color; break
            # Cajas existentes con texto directo permanecen editables.
            if not any(c.tag in ('p','h2','h3','h4','ul','ol','figure','table','div') for c in e): self.paragraph(e,fill=selected)
            else:
                if e.text and e.text.strip(): self.doc.add_paragraph(e.text.strip())
                for c in e: self.block(c,selected,level)
        elif tag not in ('script','style','nav'):
            if e.text_content().strip(): self.paragraph(e,fill=fill)
    def table(self,e):
        rows=e.xpath('./tr|./thead/tr|./tbody/tr|./tfoot/tr')
        if not rows: return
        n=max(sum(int(c.get('colspan','1')) for c in r if c.tag in ('th','td')) for r in rows)
        captions=e.xpath('./caption')
        for cap in captions:
            p=self.doc.add_paragraph(cap.text_content().strip(),style='Caption'); p.paragraph_format.keep_with_next=True
        table=self.doc.add_table(rows=0,cols=n); table.autofit=False
        # Más espacio para columnas de explicación que para etiquetas.
        sizes=[1.0]*n
        for r in rows:
            col=0
            for c in r:
                if c.tag not in ('th','td'): continue
                span=int(c.get('colspan','1'))
                sizes[col]=max(sizes[col],min(180,len(c.text_content()))**.48/span); col+=span
        total=sum(sizes)
        for col,size in zip(table.columns,sizes): col.width=Cm(16.6*size/total)
        for i,row in enumerate(rows):
            cells=table.add_row().cells; col=0; header=bool(row.xpath('./th'))
            if header:
                repeat=OxmlElement('w:tblHeader'); cells[0]._tc.getparent().get_or_add_trPr().append(repeat)
            for c in row:
                if c.tag not in ('th','td'): continue
                span=int(c.get('colspan','1')); cell=cells[col]
                if span>1: cell=cell.merge(cells[min(col+span-1,n-1)])
                cell.width=Cm(16.6*sum(sizes[col:col+span])/total)
                p=cell.paragraphs[0]; runs(p,clean(tokens(c)))
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing=1.1; p.paragraph_format.space_after=Pt(5); p.paragraph_format.space_before=Pt(4)
                    for r in p.runs:
                        r.font.size=Pt(10); r.font.bold=True if header else r.bold
                        if header: r.font.color.rgb=RGBColor(255,255,255)
                pr=cell._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'203D59' if header else ('F2F6FA' if i%2==0 else 'FFFFFF')); pr.append(sh)
                margins=OxmlElement('w:tcMar')
                for side in ('top','bottom','left','right'):
                    m=OxmlElement('w:'+side); m.set(qn('w:w'),'95'); m.set(qn('w:type'),'dxa'); margins.append(m)
                pr.append(margins); col+=span
        self.doc.add_paragraph().paragraph_format.space_after=Pt(2)
        self.stats['tablas']+=1
    def module(self,tree,number):
        article=tree.xpath('//article')[0]; title=tree.xpath('//h1')[0].text_content().strip()
        p=self.doc.add_paragraph('Gestión de Proyectos Ágiles',style='Subtitle'); p.paragraph_format.space_before=Cm(3)
        p=self.doc.add_paragraph(f'Módulo {number}'); p.runs[0].font.size=Pt(16); p.runs[0].font.color.rgb=RGBColor.from_string(PALETTE[number-1])
        self.doc.add_paragraph(title,style='Title')
        subtitle=tree.xpath('//*[contains(concat(" ",normalize-space(@class)," ")," subtitle ")]')
        if subtitle: self.doc.add_paragraph(subtitle[0].text_content().strip(),style='Subtitle')
        self.doc.add_paragraph('Contenido de estudio',style='Subtitle')
        self.doc.add_page_break()
        self.doc.add_paragraph('Recorrido del módulo',style='Heading 1')
        for i,h in enumerate(article.xpath('./h2'),1):
            p=self.doc.add_paragraph(); hyperlink(p,h.text_content().strip(),'#'+self.prefix+h.get('id',f'seccion-{i}'))
            p.paragraph_format.space_after=Pt(4)
        self.doc.add_page_break()
        for e in article: self.block(e)

def main():
    report=[]
    for n in range(1,6):
        folder=BASE/'entregables'/f'MODULO {n}'/'Contenido'; tree=html.parse(str(folder/'contenido.html'))
        doc=Document(); styles(doc,PALETTE[n-1]); builder=Builder(doc,folder); builder.module(tree,n)
        doc.core_properties.title=f'Gestión de Proyectos Ágiles Módulo {n}'
        output=folder/'Contenido.docx'; finalize(doc); doc.save(output)
        report.append({'modulo':n,'archivo':str(output),'bytes':output.stat().st_size,**builder.stats})
    # Actualizar el libro auxiliar existente para evitar una versión anterior discordante.
    doc=Document(); styles(doc,PALETTE[0]); doc.core_properties.title='Gestión de Proyectos Ágiles Contenido completo'
    for n in range(1,6):
        if n>1: doc.add_page_break()
        folder=BASE/'entregables'/f'MODULO {n}'/'Contenido'
        Builder(doc,folder,f'm{n}-').module(html.parse(str(folder/'contenido.html')),n)
    finalize(doc); doc.save(BASE/'output/GPA - materia-completa.docx')
    (BASE/'.qa/edicion-contenido/exportacion-word.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(report,ensure_ascii=False,indent=2))

if __name__=='__main__': main()
